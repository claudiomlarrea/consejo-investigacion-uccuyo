"""Word del Orden del Día y del informe por responsable, igual que el sistema actual."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any, Dict, List, Optional

from docx import Document
from docx.shared import Pt, RGBColor

from puntaje import puntaje_texto_para_word


def fila_norm(r: Dict[str, Any]) -> Dict[str, Any]:
    return {str(k).lower().strip(): v for k, v in r.items()}


def unidad_de(r: Dict[str, Any]) -> str:
    row = fila_norm(r)
    return str(row.get("unidad académica") or row.get("unidad") or "").strip()


def fingerprint(r: Dict[str, Any]) -> str:
    row = fila_norm(r)
    return "|".join(
        [
            str(row.get("tipo", "")).strip(),
            str(row.get("titulo", "") or row.get("título", "")).strip(),
            str(row.get("responsable_de_carga", "")).strip(),
            unidad_de(r),
            str(row.get("director", "")).strip(),
            str(row.get("descripcion") or row.get("descripción") or "").strip()[:120],
        ]
    )


def ordenar_por_unidad(registros: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    indexed = list(enumerate(registros))
    indexed.sort(key=lambda t: (unidad_de(t[1]).casefold(), t[0]))
    return [r for _, r in indexed]


def aplicar_orden_manual(
    registros: List[Dict[str, Any]], orden: Optional[List[str]]
) -> List[Dict[str, Any]]:
    if not registros:
        return registros
    if not orden:
        return ordenar_por_unidad(registros)
    by_fp = {}  # type: Dict[str, List[Dict[str, Any]]]
    fps = []
    for r in registros:
        fp = fingerprint(r)
        fps.append(fp)
        by_fp.setdefault(fp, []).append(r)
    usados = {}  # type: Dict[str, int]
    ordenados = []
    for fp in orden:
        cola = by_fp.get(fp) or []
        idx = usados.get(fp, 0)
        if idx < len(cola):
            ordenados.append(cola[idx])
            usados[fp] = idx + 1
    for fp in fps:
        cola = by_fp.get(fp) or []
        idx = usados.get(fp, 0)
        if idx < len(cola):
            ordenados.append(cola[idx])
            usados[fp] = idx + 1
    return ordenados


_RE_ITEM = re.compile(r"^(\d+)\.\s+(.+?)\s+-\s+(.+)$")
_RE_CAMPO = re.compile(
    r"^\s*(Descripción|Director|Codirector|Equipo|Unidad Académica|Puntaje|"
    r"Resolución CD|Resolución CS(?: del Proyecto)?|Instituto|Cátedra|Cátedras|"
    r"Financiamiento|Fuente|Responsable de carga|Monto|Alumnos|Docente|DNI):\s*(.*)$",
    re.I,
)
_MAPA_CAMPO = {
    "descripción": "descripcion",
    "director": "director",
    "codirector": "codirector",
    "equipo": "equipo",
    "unidad académica": "unidad académica",
    "puntaje": "puntaje",
    "resolución cd": "resolucion_cd",
    "resolución cs": "resolucion_cs",
    "resolución cs del proyecto": "resolucion_cs",
    "instituto": "instituto",
    "cátedra": "catedra",
    "cátedras": "catedra",
    "financiamiento": "tipo de financiamiento",
    "fuente": "fuente de financiamiento",
    "responsable de carga": "responsable_de_carga",
    "monto": "monto del financiamiento",
    "alumnos": "alumnos",
    "docente": "apellido_nombre_docente",
    "dni": "dni_docente",
}


def _aplicar_campo_od(tema: Dict[str, Any], linea: str) -> bool:
    m = _RE_CAMPO.match(linea.strip())
    if not m:
        return False
    clave = _MAPA_CAMPO.get(m.group(1).strip().casefold())
    if not clave:
        return False
    valor = m.group(2).strip()
    if clave == "director" and valor.endswith(")"):
        base, _, resto = valor.rpartition(" (")
        if base:
            tema["director"] = base.strip()
            tema["cat_director"] = resto.rstrip(")").strip()
            return True
    if clave == "codirector" and valor.endswith(")"):
        base, _, resto = valor.rpartition(" (")
        if base:
            tema["codirector"] = base.strip()
            tema["cat_codirector"] = resto.rstrip(")").strip()
            return True
    if clave == "monto del financiamiento":
        valor = valor.replace("$", "").replace(".", "").replace(" ", "")
    tema[clave] = valor
    return True


def parsear_od_docx(data: bytes, numero: int, fecha: str) -> List[Dict[str, Any]]:
    """Lee un Orden del Día generado por el sistema (el mismo formato de Acta_195.docx)."""
    doc = Document(BytesIO(data))
    temas = []  # type: List[Dict[str, Any]]
    actual = None  # type: Optional[Dict[str, Any]]
    unidad = ""
    for para in doc.paragraphs:
        lineas = [ln for ln in para.text.replace("\r", "").split("\n") if ln.strip()]
        if not lineas:
            continue
        primero = lineas[0].strip()
        item = _RE_ITEM.match(primero)
        if item:
            if actual:
                temas.append(actual)
            actual = {
                "numero_acta": numero,
                "fecha": fecha,
                "año": 2026,
                "tipo": item.group(2).strip(),
                "titulo": item.group(3).strip(),
                "unidad académica": unidad,
                "_fuente": "word",
            }
            for extra in lineas[1:]:
                _aplicar_campo_od(actual, extra)
            continue
        if actual:
            if not _aplicar_campo_od(actual, primero):
                pass
            continue
        if primero not in {"Consejo de Investigación", "Orden del Día"} and not primero.startswith("Acta N") and not primero.startswith("Fecha:"):
            unidad = primero
    if actual:
        temas.append(actual)
    return temas


def construir_od(acta_num: int, registros: List[Dict[str, Any]]) -> bytes:
    doc = Document()
    doc.add_heading("Consejo de Investigación", 0)
    p_acta = doc.add_paragraph("Acta N° {}".format(acta_num))
    p_acta.paragraph_format.space_after = Pt(0)
    fecha_real = ""
    if registros:
        first = fila_norm(registros[0])
        fecha_real = first.get("fecha") or first.get("FECHA") or ""
    p_fecha = doc.add_paragraph("Fecha: {}".format(fecha_real))
    p_fecha.paragraph_format.space_after = Pt(0)
    doc.add_heading("Orden del Día", level=1)

    contador = 1
    unidad_actual = None
    for r in registros:
        r = fila_norm(r)
        unidad = str(r.get("unidad académica") or r.get("unidad") or "").strip()
        if unidad != unidad_actual:
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.line_spacing = 1
            run_h = h.add_run(unidad)
            run_h.bold = True
            run_h.font.color.rgb = RGBColor(0, 102, 204)
            unidad_actual = unidad

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1
        p.add_run("{}. {} - {}\n".format(contador, r.get("tipo", ""), r.get("titulo", ""))).bold = True

        descripcion = r.get("descripcion") or r.get("descripción") or ""
        if descripcion:
            p.add_run("   Descripción: {}\n".format(descripcion))

        tipo = r.get("tipo", "")
        if tipo == "Categorización Docente":
            if r.get("apellido_nombre_docente"):
                p.add_run("   Docente: {}\n".format(r.get("apellido_nombre_docente")))
            if r.get("dni_docente"):
                p.add_run("   DNI: {}\n".format(r.get("dni_docente")))

        if tipo in (
            "Proyecto de Investigación",
            "Proyecto de Cátedra",
            "Informe Final",
            "Informe de Avance",
        ):
            cat = r.get("cat_director", "")
            if cat in ("Seleccionar", "", None):
                p.add_run("   Director: {}\n".format(r.get("director", "")))
            else:
                p.add_run("   Director: {} ({})\n".format(r.get("director", ""), cat))
            cat_c = r.get("cat_codirector", "")
            if cat_c in ("Seleccionar", "", None):
                p.add_run("   Codirector: {}\n".format(r.get("codirector", "")))
            else:
                p.add_run("   Codirector: {} ({})\n".format(r.get("codirector", ""), cat_c))

        equipo = r.get("equipo") or ""
        if equipo:
            p.add_run("   Equipo: {}\n".format(str(equipo).replace("\n", "; ")))
        p.add_run("   Unidad Académica: {}\n".format(unidad))
        txt_p = puntaje_texto_para_word(r.get("puntaje"))
        if txt_p:
            p.add_run("   Puntaje: {}\n".format(txt_p))
        if r.get("resolucion_cd"):
            p.add_run("   Resolución CD: {}\n".format(r.get("resolucion_cd")))
        if r.get("resolucion_cs"):
            p.add_run("   Resolución CS del Proyecto: {}\n".format(r.get("resolucion_cs")))
        if r.get("instituto"):
            p.add_run("   Instituto: {}\n".format(r.get("instituto")))
        catedra = r.get("catedra") or r.get("cátedra") or ""
        if str(catedra).strip():
            etiqueta = "Cátedras" if (";" in str(catedra) or "," in str(catedra)) else "Cátedra"
            p.add_run("   {}: {}\n".format(etiqueta, catedra))
        if r.get("tipo de financiamiento"):
            p.add_run("   Financiamiento: {}\n".format(r.get("tipo de financiamiento")))
        if r.get("fuente de financiamiento"):
            p.add_run("   Fuente: {}\n".format(r.get("fuente de financiamiento")))
        if r.get("responsable_de_carga"):
            p.add_run("   Responsable de carga: {}\n".format(r.get("responsable_de_carga")))
        if r.get("monto del financiamiento"):
            try:
                monto = int(float(r.get("monto del financiamiento")))
                monto = "${}".format("{:,}".format(monto).replace(",", "."))
            except Exception:
                monto = r.get("monto del financiamiento")
            p.add_run("   Monto: {}\n".format(monto))
        if r.get("alumnos"):
            p.add_run("   Alumnos: {}\n".format(r.get("alumnos")))
        contador += 1

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def construir_informe(acta_num: int, responsable: str, registros: List[Dict[str, Any]]) -> bytes:
    doc = Document()
    doc.add_heading("Informe de temas cargados", 0)
    doc.add_paragraph("Acta N° {}".format(acta_num))
    doc.add_paragraph("Responsable de carga: {}".format(responsable))
    contador = 1
    unidad_actual = None
    for r in registros:
        r = fila_norm(r)
        unidad = str(r.get("unidad académica") or r.get("unidad") or "").strip()
        if unidad != unidad_actual:
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            run_h = h.add_run(unidad)
            run_h.bold = True
            run_h.font.color.rgb = RGBColor(0, 102, 204)
            unidad_actual = unidad
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.add_run("{}. {} - {}\n".format(contador, r.get("tipo", ""), r.get("titulo", ""))).bold = True
        descripcion = r.get("descripcion") or r.get("descripción") or ""
        if descripcion:
            p.add_run("   Descripción: {}\n".format(descripcion))
        p.add_run("   Unidad Académica: {}\n".format(unidad))
        contador += 1
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
